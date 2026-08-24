from __future__ import annotations

import argparse
import hashlib
import json
import re
import unicodedata
from collections import Counter
from pathlib import Path

from docx import Document
from pypdf import PdfReader


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def normalize(text: str) -> list[str]:
    text = unicodedata.normalize("NFKC", text).casefold()
    return re.findall(r"[\w]+", text, flags=re.UNICODE)


def multiset_similarity(left: list[str], right: list[str]) -> float:
    a, b = Counter(left), Counter(right)
    intersection = sum((a & b).values())
    union = sum((a | b).values())
    return intersection / union if union else 1.0


def docx_text(path: Path) -> tuple[str, dict[str, object]]:
    document = Document(path)
    parts: list[str] = []
    nonempty_paragraphs = 0
    headings: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            nonempty_paragraphs += 1
            parts.append(text)
            if paragraph.style and paragraph.style.name.startswith("Heading"):
                headings.append(text)
    table_cells = 0
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                value = cell.text.strip()
                if value:
                    table_cells += 1
                    parts.append(value)
    core = document.core_properties
    meta = {
        "paragraphs_total": len(document.paragraphs),
        "paragraphs_nonempty": nonempty_paragraphs,
        "tables": len(document.tables),
        "nonempty_table_cells": table_cells,
        "headings_count": len(headings),
        "headings": headings,
        "core_properties": {
            "title": core.title,
            "subject": core.subject,
            "author": core.author,
            "last_modified_by": core.last_modified_by,
            "created": core.created.isoformat() if core.created else None,
            "modified": core.modified.isoformat() if core.modified else None,
        },
    }
    return "\n".join(parts), meta


def markdown_text(path: Path) -> str:
    return path.read_text(encoding="utf-8-sig")


def pdf_text(path: Path) -> tuple[str, dict[str, object]]:
    reader = PdfReader(path)
    page_text = [(page.extract_text() or "") for page in reader.pages]
    return "\n".join(page_text), {
        "pages": len(reader.pages),
        "pages_with_text": sum(bool(text.strip()) for text in page_text),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--markdown", type=Path, required=True)
    parser.add_argument("--docx", type=Path, required=True)
    parser.add_argument("--pdf", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    md = markdown_text(args.markdown)
    docx, docx_meta = docx_text(args.docx)
    pdf, pdf_meta = pdf_text(args.pdf)
    tokens = {"markdown": normalize(md), "docx": normalize(docx), "pdf": normalize(pdf)}
    report = {
        "schema": "bayme-eve-canonical-representation-adjudication/v1",
        "classification": "DERIVATIVE_REPRESENTATIONS_OF_ONE_CANONICAL_TEXT",
        "authority": "semantic-and-byte-custody-only",
        "visual_gate": "NOT_PASSED_VIEWER_ACL_BLOCKED",
        "files": {
            "markdown": {"path": str(args.markdown), "bytes": args.markdown.stat().st_size, "sha256": sha256(args.markdown), "tokens": len(tokens["markdown"])},
            "docx": {"path": str(args.docx), "bytes": args.docx.stat().st_size, "sha256": sha256(args.docx), "tokens": len(tokens["docx"]), **docx_meta},
            "pdf": {"path": str(args.pdf), "bytes": args.pdf.stat().st_size, "sha256": sha256(args.pdf), "tokens": len(tokens["pdf"]), **pdf_meta},
        },
        "token_multiset_jaccard": {
            "markdown_docx": multiset_similarity(tokens["markdown"], tokens["docx"]),
            "markdown_pdf": multiset_similarity(tokens["markdown"], tokens["pdf"]),
            "docx_pdf": multiset_similarity(tokens["docx"], tokens["pdf"]),
        },
        "adjudication": {
            "scientific_unit": "hbp-economic-political-core",
            "new_market_unit": False,
            "bum_unit": False,
            "bpm_unit": False,
            "claims": "Constitutional and software-alpha specification; not an executed validation of BayME-EVE as a complete system.",
        },
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
